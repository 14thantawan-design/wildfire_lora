"""Build a Thai DOCX reference for the Sensor Node threshold model (v8).

The document is generated from the fixed thresholds used by NODE01/NODE02.
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "แหล่งอ้างอิงหลักการคำนวณไฟป่า_ฝั่งโหนด.docx"
FONT = "Tahoma"
NAVY = "17324D"
BLUE = "2F6F9F"
LIGHT_BLUE = "EAF3F8"
LIGHT_YELLOW = "FFF6DA"
LIGHT_RED = "FDE8E5"


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, *, bold: bool = False, color: str = "222222") -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = FONT
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, bold=True, color="FFFFFF")
        set_cell_shading(table.rows[0].cells[index], BLUE)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_text(cells[index], value)
    document.add_paragraph()


def add_callout(document: Document, title: str, body: str, fill: str) -> None:
    table = document.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    paragraph = cell.paragraphs[0]
    title_run = paragraph.add_run(f"{title} ")
    title_run.bold = True
    body_run = paragraph.add_run(body)
    for run in (title_run, body_run):
        run.font.name = FONT
        run.font.size = Pt(10)
    document.add_paragraph()


def add_bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.add_run(text)


def source_line(relative_path: str, marker: str) -> str:
    path = ROOT / relative_path
    for line_number, line in enumerate(path.read_text(encoding="utf-8").splitlines(), start=1):
        if marker in line:
            return f"{relative_path}:{line_number}"
    return f"{relative_path} (ค้นด้วยคำว่า {marker})"


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    normal = document.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.15

    for style_name, size in (("Title", 22), ("Heading 1", 15), ("Heading 2", 12)):
        style = document.styles[style_name]
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string("000000")
        if style_name == "Title":
            style_properties = style.element.get_or_add_pPr()
            style_border = style_properties.find(qn("w:pBdr"))
            if style_border is not None:
                style_properties.remove(style_border)


def build_document() -> Path:
    document = Document()
    configure_document(document)

    title = document.add_heading("หลักการคำนวณสถานะไฟป่าฝั่ง Sensor Node", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_properties = title._p.get_or_add_pPr()
    title_border = title_properties.find(qn("w:pBdr"))
    if title_border is not None:
        title_properties.remove(title_border)
    subtitle = document.add_paragraph("Risk model version 8 · NODE01 และ NODE02")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph(f"ปรับปรุงเอกสาร: {date.today().isoformat()}").alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_callout(
        document,
        "ขอบเขต:",
        "ใช้สถานะ NORMAL, WATCH และ WARNING เท่านั้น ส่วน SENSOR_FAULT เป็นสุขภาพอุปกรณ์ "
        "ไม่มีคะแนน 0–100 ไม่มีค่าฐาน และไม่มีอัตราการเปลี่ยนแปลงในการตัดสิน",
        LIGHT_BLUE,
    )

    document.add_heading("1. ค่าที่ระบบใช้", level=1)
    add_table(
        document,
        ["ตัวแปร", "แหล่งวัด", "หน่วย/ความหมาย"],
        [
            ["T", "SHT31", "อุณหภูมิอากาศ (°C)"],
            ["RH", "SHT31", "ความชื้นสัมพัทธ์ (%RH)"],
            ["P", "Sharp GP2Y1014AU0F", "ความเข้มข้นอนุภาคโดยประมาณ (µg/m³)"],
        ],
    )
    add_callout(
        document,
        "ข้อจำกัดของ P:",
        "GP2Y1014AU0F ไม่ได้แยกขนาดอนุภาค จึงห้ามเรียกผลนี้ว่า PM2.5 ที่ผ่านการสอบเทียบ "
        "ค่าต้องผ่านการทดลองเทียบเครื่องอ้างอิงก่อนกล่าวอ้างความแม่นยำ",
        LIGHT_YELLOW,
    )

    document.add_heading("2. การแปลงค่า Sharp", level=1)
    document.add_paragraph("อ่านแรงดันสามครั้ง เลือกค่ามัธยฐาน แล้วคำนวณค่าประมาณดังนี้")
    equation = document.add_paragraph("P = max(0, (Vo_mV - 600) / 5)")
    equation.alignment = WD_ALIGN_PARAGRAPH.CENTER
    equation.runs[0].bold = True
    add_bullet(document, "600 mV คือค่าแรงดันศูนย์ทั่วไปที่ตั้งต้นไว้และควรวัดใหม่แยกแต่ละอุปกรณ์")
    add_bullet(document, "5 mV ต่อ µg/m³ มาจากความไวทั่วไป 0.5 V ต่อ 100 µg/m³ ในเอกสารผู้ผลิต")
    add_bullet(document, "เกณฑ์ 50/150 เป็นการดัดแปลงมาใช้กับค่าประมาณของ Sharp ไม่ใช่การแทนค่า MQ-2 raw")

    document.add_heading("3. ลำดับการตัดสินสถานะ", level=1)
    rules = document.add_paragraph()
    rules.add_run(
        "ถ้าเซนเซอร์ผิดปกติ → SENSOR_FAULT\n"
        "มิฉะนั้น ถ้า T > 45 หรือ P > 150 หรือ (T ≥ 30 และ RH ≤ 30) → WARNING\n"
        "มิฉะนั้น ถ้า T > 35 หรือ P > 50 หรือ RH < 50 → WATCH\n"
        "มิฉะนั้น → NORMAL"
    )
    add_table(
        document,
        ["สถานะ", "เงื่อนไข", "รอบวัดและส่ง"],
        [
            ["NORMAL", "T ≤ 35 และ P ≤ 50 และ RH ≥ 50", "300 วินาที"],
            ["WATCH", "ยังไม่เป็น WARNING และ T > 35 หรือ P > 50 หรือ RH < 50", "120 วินาที"],
            ["WARNING", "T > 45 หรือ P > 150 หรือ (T ≥ 30 และ RH ≤ 30)", "ส่งทันที แล้วทุก 20 วินาที"],
            ["SENSOR_FAULT", "SHT31 หรือ Sharp ให้ข้อมูลใช้ไม่ได้", "300 วินาที"],
        ],
    )

    document.add_heading("4. บทบาทของความชื้น", level=1)
    add_bullet(document, "RH < 50% เพียงอย่างเดียวทำให้เป็น WATCH")
    add_bullet(document, "RH ≤ 30% ร่วมกับ T ≥ 30°C ทำให้เป็น WARNING")
    add_bullet(document, "RH ≤ 30% แต่ T < 30°C ยังคงเป็น WATCH")
    add_callout(
        document,
        "ที่มาของกฎร้อน–แห้ง:",
        "ประยุกต์สองปัจจัยที่วัดได้จากแนวคิด 30-30-30 ใน designs-09-00091-v3.pdf "
        "ระบบไม่มีเซนเซอร์ลม จึงไม่กล่าวอ้างว่าใช้กฎครบสามปัจจัย",
        LIGHT_BLUE,
    )

    document.add_heading("5. การยืนยันก่อนลดระดับ", level=1)
    add_bullet(document, "ยกระดับทันทีเมื่อค่ารอบปัจจุบันผ่านเงื่อนไข")
    add_bullet(document, "WARNING ลดเป็น WATCH หลังไม่เข้า WARNING ติดต่อกัน 3 รอบ")
    add_bullet(document, "WATCH ลดเป็น NORMAL หลังเข้า NORMAL ติดต่อกัน 3 รอบ")
    add_bullet(document, "WARNING ที่กลับปกติจะผ่าน WATCH ก่อนเสมอ")

    document.add_heading("6. ข้อมูลที่ส่งในแพ็กเก็ต", level=1)
    add_bullet(document, "โหนดส่งค่า T, RH และ P ที่วัดได้")
    add_bullet(document, "โหนดส่งสถานะสุดท้าย NORMAL, WATCH, WARNING หรือ SENSOR_FAULT")
    add_bullet(document, "Gateway และ Backend ไม่คำนวณเกณฑ์ซ้ำและไม่มีรหัสเหตุผลแบบบิต")

    document.add_heading("7. การเชื่อมโยงแหล่งอ้างอิง", level=1)
    add_table(
        document,
        ["แหล่ง", "สิ่งที่นำมาใช้", "สถานะการนำมาใช้"],
        [
            ["computers-15-00105.pdf", "ช่วงอุณหภูมิ 35/45°C และช่วงอนุภาค 50/150", "ดัดแปลงชนิดเซนเซอร์และชื่อสถานะให้ตรงโครงการ"],
            ["designs-09-00091-v3.pdf", "อุณหภูมิอย่างน้อย 30°C ร่วม RH ไม่เกิน 30%", "ใช้ 2 ใน 3 ปัจจัยเพราะไม่มีเซนเซอร์ลม"],
            ["Sharp GP2Y1014AU0F specification", "จังหวะ LED และความไวทั่วไปของเอาต์พุต", "ใช้สร้างค่าประมาณ ต้องสอบเทียบหน้างาน"],
        ],
    )
    add_callout(
        document,
        "คำที่ควรใช้ในรายงาน:",
        "ระบบต้นแบบตรวจสัญญาณบ่งชี้จากอนุภาค อุณหภูมิ และความชื้น "
        "ไม่ควรอ้างว่าตรวจหรือทำนายไฟป่าได้แน่นอน 100%",
        LIGHT_RED,
    )

    document.add_page_break()
    document.add_heading("8. ตำแหน่งใน Source code", level=1)
    add_table(
        document,
        ["หัวข้อ", "ตำแหน่งค้นสอบ"],
        [
            ["แปลงค่า Sharp", source_line("sensor_node/sensor_node.ino", "particleUgM3FromMilliVolts")],
            ["ตัดสินค่าดิบ", source_line("sensor_node/sensor_node.ino", "FireStatus evaluateRawRisk")],
            ["ยืนยันการลดระดับ", source_line("sensor_node/sensor_node.ino", "FireStatus applyStateLatch")],
            ["สร้างแพ็กเก็ต", source_line("sensor_node/sensor_node.ino", "String buildJsonPacket")],
            ["ค่าคงที่", source_line("sensor_node/config.h", "RISK_MODEL_VERSION")],
            ["ชุดทดสอบ", "tools/test_node_risk.mjs"],
        ],
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(f"DOCX_CREATED: {OUTPUT}")
    return OUTPUT


if __name__ == "__main__":
    build_document()
